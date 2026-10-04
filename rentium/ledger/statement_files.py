"""Bounded, read-only bank export parsing. No formula evaluation or guessed dates."""

import csv
import io
import zipfile
from itertools import islice
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%m/%d/%y", "%d/%m/%y")
FIELDS = ("date", "description", "amount", "debit", "credit", "reference")


def read_statement(content, filename, worksheet=""):
    if len(content) > MAX_BYTES:
        raise ValueError("Choose a statement smaller than 10 MB.")
    sheets = []
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("The expanded workbook is too large.")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            sheets = workbook.sheetnames
            selected = worksheet
            if not selected:
                for name in sheets:
                    first = list(islice(workbook[name].iter_rows(values_only=True), 2))
                    if len(first) == 2 and len(first[0]) > 1 and all(first[0]):
                        selected = name
                        break
                selected = selected or sheets[0]
            if selected not in sheets:
                raise ValueError("Choose one of the worksheets in this file.")
            sheet = workbook[selected]
            if sheet.max_column and sheet.max_column > 64:
                raise ValueError(
                    "Choose a tabular bank export with at most 64 columns."
                )
            matrix = []
            for row in sheet.iter_rows(values_only=True):
                if len(matrix) > MAX_ROWS:
                    raise ValueError(
                        "Split the statement into files of at most 5,000 rows."
                    )
                if any(isinstance(v, str) and v.startswith("=") for v in row):
                    raise ValueError(
                        "Bank exports must contain values, not Excel formulas."
                    )
                matrix.append(
                    [
                        v.isoformat()
                        if isinstance(v, (datetime, date))
                        else str(v)
                        if v is not None
                        else ""
                        for v in row
                    ]
                )
        finally:
            workbook.close()
    elif filename.lower().endswith(".csv"):
        selected = ""
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        matrix = []
        for row in reader:
            if len(matrix) > MAX_ROWS or len(row) > 64:
                raise ValueError("Use at most 5,000 rows and 64 columns per statement.")
            matrix.append(row)
    else:
        raise ValueError(
            "Use a CSV or XLSX bank export. Upload receipt PDFs/photos in Documents."
        )
    matrix = [row for row in matrix if any(str(v).strip() for v in row)]
    if len(matrix) < 2:
        raise ValueError(
            "The first row must contain column names, followed by transactions."
        )
    headers = [str(v).strip() for v in matrix[0]]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError("Every column needs a unique, nonempty name in the first row.")
    if any(len(row) != len(headers) for row in matrix[1:]):
        raise ValueError(
            "Every transaction must have the same number of columns as the header."
        )
    return headers, [dict(zip(headers, row)) for row in matrix[1:]], sheets, selected


def guess_mapping(headers):
    aliases = {
        "date": "date",
        "transaction date": "date",
        "posted date": "date",
        "description": "description",
        "memo": "description",
        "details": "description",
        "amount": "amount",
        "debit": "debit",
        "withdrawals": "debit",
        "withdrawal": "debit",
        "credit": "credit",
        "deposits": "credit",
        "deposit": "credit",
        "reference": "reference",
        "transaction id": "reference",
    }
    return {
        header: aliases[header.lower()]
        for header in headers
        if header.lower() in aliases
    }


def decimal_amount(raw):
    text = str(raw).strip().replace(",", "").replace("$", "")
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("Amount is missing or invalid.") from exc
    if (
        not value.is_finite()
        or abs(value) >= Decimal("1000000000000")
        or value != value.quantize(Decimal("0.01"))
    ):
        raise ValueError("Amounts must be finite and have at most two decimal places.")
    return value


def parse_row(raw, mapping, date_format):
    if date_format not in DATE_FORMATS:
        raise ValueError("Choose a supported date format.")
    values = {
        target: raw.get(source, "") for source, target in mapping.items() if target
    }
    text_date = values.get("date", "").strip()
    try:
        booked = datetime.strptime(text_date, date_format).date()
    except ValueError:
        # Excel typed dates are exported by the reader in ISO format.
        try:
            booked = datetime.fromisoformat(text_date).date()
        except ValueError as exc:
            raise ValueError("Date does not match the selected date format.") from exc
    if "amount" in values:
        amount = decimal_amount(values["amount"])
    else:
        debit = decimal_amount(values.get("debit") or "0")
        credit = decimal_amount(values.get("credit") or "0")
        if debit < 0 or credit < 0 or (debit and credit):
            raise ValueError(
                "Use nonnegative debit/credit columns, with only one populated per row."
            )
        amount = credit - debit
    if not amount:
        raise ValueError("Zero-value rows are not bank transactions.")
    return (
        booked,
        amount,
        values.get("description", "").strip(),
        values.get("reference", "").strip(),
    )
